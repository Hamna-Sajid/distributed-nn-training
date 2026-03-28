"""
benchmark_suite.py — Automated benchmarking across multiple sample sizes.

Reads experiment parameters from config.yaml (benchmark.sample_sizes,
benchmark.n_epochs_per_run). For each sample size it:
  1. Spawns a master process and two worker processes
  2. Runs full distributed training
  3. Saves a timestamped JSON file per run
  4. Appends a row to benchmark_results/benchmark_summary.csv
  5. Rebuilds benchmark_results/benchmark_summary.xlsx after all runs

Usage
-----
    python benchmark_suite.py

You do NOT need to open separate terminals. The script manages
all processes internally.

Output
------
    benchmark_results/
        run_YYYYMMDD_HHMMSS_Nsamples.json   -- raw metrics per run
        benchmark_summary.csv               -- master CSV (one row per run)
        benchmark_summary.xlsx              -- formatted Excel workbook
"""

import os
import sys
import json
import time
import datetime
import csv
import multiprocessing

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)

from config_loader import CFG


# ─────────────────────────────────────────────────────────────────────────────
# Process target functions
# ─────────────────────────────────────────────────────────────────────────────

def _master_target(n_samples: int, n_epochs: int, port: int,
                   result_queue: multiprocessing.Queue):
    """Run the master in a subprocess and post results to result_queue.

    Parameters
    ----------
    n_samples : int
        MNIST samples to use for this experiment.
    n_epochs : int
        Training epochs for this experiment.
    port : int
        TCP port to bind on.
    result_queue : multiprocessing.Queue
        Queue where the master posts its result dict when training finishes.
    """
    try:
        from coordination.master import Master
        master = Master(
            n_workers=CFG["training"]["n_workers"],
            host=CFG["communication"]["host"],
            port=port,
            n_epochs=n_epochs,
            lr=CFG["model"]["lr"],
            n_samples=n_samples,
        )
        metrics = master.run_and_return_metrics()
        result_queue.put({"status": "ok", "metrics": metrics})
    except Exception as exc:
        result_queue.put({"status": "error", "error": str(exc)})


def _worker_target(worker_id: int, port: int, delay: float = 0.0):
    """Run a worker in a subprocess.

    Parameters
    ----------
    worker_id : int
        Numeric ID (0 = fast worker, 1 = slow worker).
    port : int
        TCP port of the master to connect to.
    delay : float
        Artificial delay in seconds to simulate a slow node.
    """
    host = CFG["communication"]["host"]
    from coordination.worker import Worker

    # Retry connection because process scheduling can start workers
    # before the master has finished binding/listening on the port.
    max_attempts = 30
    for attempt in range(1, max_attempts + 1):
        try:
            w = Worker(
                worker_id=worker_id,
                master_host=host,
                master_port=port,
                artificial_delay=delay,
            )
            w.run()
            return
        except ConnectionRefusedError:
            if attempt == max_attempts:
                print(
                    f"[Benchmark] Worker {worker_id} failed to connect "
                    f"to {host}:{port} after {max_attempts} attempts."
                )
                raise
            time.sleep(0.5)



# ─────────────────────────────────────────────────────────────────────────────
# Single experiment
# ─────────────────────────────────────────────────────────────────────────────

def run_experiment(n_samples: int, n_epochs: int, port: int) -> dict:
    """Run one complete distributed training experiment.

    Spawns master + 2 workers as separate processes, waits for master
    to finish, collects results from the queue.

    Parameters
    ----------
    n_samples : int
        MNIST samples to use.
    n_epochs : int
        Training epochs.
    port : int
        TCP port for this experiment (unique per run to avoid conflicts).

    Returns
    -------
    dict
        All metrics plus n_samples, n_epochs, total_time_sec, status.
    """
    print(f"\n{'='*58}")
    print(f"  EXPERIMENT  |  {n_samples:,} samples  |  {n_epochs} epochs  |  port {port}")
    print(f"{'='*58}")

    result_queue = multiprocessing.Queue()

    master_proc  = multiprocessing.Process(
        target=_master_target,
        args=(n_samples, n_epochs, port, result_queue),
        daemon=False,
    )
    worker0_proc = multiprocessing.Process(
        target=_worker_target,
        args=(0, port, 0.0),   # fast worker
        daemon=True,
    )
    worker1_proc = multiprocessing.Process(
        target=_worker_target,
        args=(1, port, 2.0),   # slow worker — 2s delay
        daemon=True,
    )

    t_start = time.time()
    master_proc.start()

    worker0_proc.start()
    worker1_proc.start()

    timeout = CFG["communication"]["timeout_sec"] * n_epochs
    master_proc.join(timeout=timeout)
    total_time = round(time.time() - t_start, 2)

    if master_proc.is_alive():
        master_proc.terminate()
        print(f"[Benchmark] TIMEOUT after {total_time}s")
        return {
            "n_samples": n_samples, "n_epochs": n_epochs,
            "total_time_sec": total_time, "status": "timeout",
        }

    if result_queue.empty():
        return {
            "n_samples": n_samples, "n_epochs": n_epochs,
            "total_time_sec": total_time, "status": "error",
            "error": "master exited without posting results",
        }

    result = result_queue.get()
    if result["status"] != "ok":
        return {
            "n_samples": n_samples, "n_epochs": n_epochs,
            "total_time_sec": total_time,
            "status": result.get("status", "error"),
            "error": result.get("error", "unknown"),
        }

    m = result["metrics"]
    return {
        "n_samples":              n_samples,
        "n_epochs":               n_epochs,
        "total_time_sec":         total_time,
        "avg_time_per_epoch_sec": round(total_time / n_epochs, 2),
        "final_loss":             m.get("final_loss"),
        "test_loss":              m.get("test_loss"),
        "accuracy":               m.get("accuracy"),
        "f1":                     m.get("f1"),
        "precision":              m.get("precision"),
        "recall":                 m.get("recall"),
        "worker0_final_loss":     m.get("worker0_final_loss"),
        "worker1_final_loss":     m.get("worker1_final_loss"),
        "worker0_samples":        m.get("worker0_samples"),
        "worker1_samples":        m.get("worker1_samples"),
        "worker0_rtt":            m.get("worker0_rtt"),
        "worker1_rtt":            m.get("worker1_rtt"),
        "epoch_times":            m.get("epoch_times", []),
        "status":                 "ok",
    }


# ─────────────────────────────────────────────────────────────────────────────
# Storage helpers
# ─────────────────────────────────────────────────────────────────────────────

def save_json(result: dict, results_dir: str) -> str:
    """Save one run's results as a timestamped JSON file.

    Parameters
    ----------
    result : dict
        Result dictionary from run_experiment().
    results_dir : str
        Directory to save into.

    Returns
    -------
    str
        Full path to the saved JSON file.
    """
    os.makedirs(results_dir, exist_ok=True)
    ts    = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    fname = f"run_{ts}_{result['n_samples']}samples.json"
    fpath = os.path.join(results_dir, fname)
    with open(fpath, "w") as f:
        json.dump(result, f, indent=2)
    print(f"[Benchmark] JSON saved -> {fpath}")
    return fpath


def append_csv(result: dict, results_dir: str):
    """Append one result row to the master CSV file.

    Creates the CSV with a header row if it does not yet exist.
    Safe to call after every run — never overwrites previous rows.

    Parameters
    ----------
    result : dict
        Result dictionary from run_experiment().
    results_dir : str
        Directory containing benchmark_summary.csv.
    """
    os.makedirs(results_dir, exist_ok=True)
    fpath  = os.path.join(results_dir, "benchmark_summary.csv")
    fields = [
        "timestamp", "n_samples", "n_epochs",
        "total_time_sec", "avg_time_per_epoch_sec",
        "final_loss", "test_loss", "accuracy", "f1",
        "precision", "recall",
        "worker0_final_loss", "worker1_final_loss",
        "worker0_samples", "worker1_samples",
        "worker0_rtt", "worker1_rtt", "status",
    ]
    write_header = not os.path.exists(fpath)
    with open(fpath, "a", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        if write_header:
            writer.writeheader()
        row = dict(result)
        row["timestamp"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        row.pop("epoch_times", None)
        writer.writerow(row)
    print(f"[Benchmark] CSV updated -> {fpath}")


# ─────────────────────────────────────────────────────────────────────────────
# Excel report builder
# ─────────────────────────────────────────────────────────────────────────────

def build_excel(results_dir: str):
    """Build a formatted Excel workbook from all run JSON files.

    Creates three sheets:
      Summary   -- one row per experiment, all metrics, colour-banded rows
      ChartData -- clean tables for each graph (time, loss, accuracy)
      RunLog    -- timestamped list of every JSON file saved

    Parameters
    ----------
    results_dir : str
        Directory containing run_*.json files.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[Benchmark] openpyxl not installed -- skipping Excel.\n"
              "            pip install openpyxl")
        return

    runs = []
    for fname in sorted(os.listdir(results_dir)):
        if fname.startswith("run_") and fname.endswith(".json"):
            with open(os.path.join(results_dir, fname)) as f:
                runs.append(json.load(f))
    if not runs:
        print("[Benchmark] No run JSON files found.")
        return
    runs.sort(key=lambda r: r.get("n_samples", 0))

    C_HDR1  = "1F4E79"
    C_HDR2  = "2E75B6"
    C_BAND1 = "FFFFFF"
    C_BAND2 = "EBF3FB"
    C_TEXT  = "595959"
    C_WHITE = "FFFFFF"
    thin    = Side(style="thin", color="CCCCCC")
    border  = Border(left=thin, right=thin, top=thin, bottom=thin)

    def title_row(ws, text, ncols, row=1):
        end_col = get_column_letter(ncols)
        ws.merge_cells(f"A{row}:{end_col}{row}")
        c = ws[f"A{row}"]
        c.value     = text
        c.font      = Font(bold=True, size=13, color=C_WHITE, name="Arial")
        c.fill      = PatternFill("solid", fgColor=C_HDR1)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 22

    def hdr(ws, row, col, text):
        c = ws.cell(row=row, column=col, value=text)
        c.font      = Font(bold=True, size=10, color=C_WHITE, name="Arial")
        c.fill      = PatternFill("solid", fgColor=C_HDR2)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border    = border
        ws.row_dimensions[row].height = 30

    def dat(ws, row, col, value, fmt=None, is_odd=True):
        c = ws.cell(row=row, column=col, value=value)
        c.font      = Font(size=10, color=C_TEXT, name="Arial")
        c.fill      = PatternFill("solid",
                                  fgColor=C_BAND1 if is_odd else C_BAND2)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = border
        if fmt:
            c.number_format = fmt

    wb = Workbook()

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.sheet_view.showGridLines = False

    COLS = [
        ("Samples",         "n_samples",              "#,##0",  11),
        ("Epochs",          "n_epochs",               "#,##0",  8),
        ("Total Time (s)",  "total_time_sec",          "0.00",   16),
        ("Time/Epoch (s)",  "avg_time_per_epoch_sec",  "0.00",   16),
        ("Final Loss",      "final_loss",              "0.0000", 12),
        ("Test Loss",       "test_loss",               "0.0000", 12),
        ("Accuracy",        "accuracy",                "0.0000", 12),
        ("F1 Score",        "f1",                      "0.0000", 12),
        ("W0 Loss",         "worker0_final_loss",       "0.0000", 12),
        ("W1 Loss",         "worker1_final_loss",       "0.0000", 12),
        ("W0 Samples",      "worker0_samples",          "#,##0",  12),
        ("W1 Samples",      "worker1_samples",          "#,##0",  12),
        ("W0 RTT (s)",      "worker0_rtt",             "0.0000", 12),
        ("W1 RTT (s)",      "worker1_rtt",             "0.0000", 12),
        ("Status",          "status",                  None,     10),
        ("Timestamp",       None,                      None,     20),
    ]

    title_row(ws1, "Distributed Neural Network Training — Benchmark Summary",
              len(COLS))
    for ci, (label, _, fmt, width) in enumerate(COLS, 1):
        hdr(ws1, 2, ci, label)
        ws1.column_dimensions[get_column_letter(ci)].width = width

    ts_now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    for ri, run in enumerate(runs, 3):
        odd = (ri % 2 == 1)
        for ci, (_, key, fmt, _) in enumerate(COLS[:-1], 1):
            dat(ws1, ri, ci, run.get(key), fmt=fmt, is_odd=odd)
        dat(ws1, ri, len(COLS), ts_now, is_odd=odd)

    ws1.freeze_panes = "A3"

    # ── ChartData sheet ───────────────────────────────────────────────────────
    ws2 = wb.create_sheet("ChartData")
    ws2.sheet_view.showGridLines = False

    SECTIONS = [
        ("Training Time vs Sample Size", [
            ("Sample Size",    "n_samples",             "#,##0",  14),
            ("Total Time (s)", "total_time_sec",         "0.00",   14),
            ("Time/Epoch (s)", "avg_time_per_epoch_sec", "0.00",   14),
        ]),
        ("Loss vs Sample Size", [
            ("Sample Size", "n_samples",  "#,##0",  14),
            ("Final Loss",  "final_loss", "0.0000", 14),
            ("Test Loss",   "test_loss",  "0.0000", 14),
        ]),
        ("Accuracy vs Sample Size", [
            ("Sample Size", "n_samples", "#,##0",  14),
            ("Accuracy",    "accuracy",  "0.0000", 14),
            ("F1 Score",    "f1",        "0.0000", 14),
        ]),
    ]

    current_col = 1
    for section_title, columns in SECTIONS:
        ncols            = len(columns)
        start_ltr        = get_column_letter(current_col)
        end_ltr          = get_column_letter(current_col + ncols - 1)
        ws2.merge_cells(f"{start_ltr}1:{end_ltr}1")
        c            = ws2[f"{start_ltr}1"]
        c.value      = section_title
        c.font       = Font(bold=True, size=11, color=C_WHITE, name="Arial")
        c.fill       = PatternFill("solid", fgColor=C_HDR1)
        c.alignment  = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 22
        for ci, (label, _, _, width) in enumerate(columns):
            abs_col = current_col + ci
            hdr(ws2, 2, abs_col, label)
            ws2.column_dimensions[get_column_letter(abs_col)].width = width
        for ri, run in enumerate(runs, 3):
            odd = (ri % 2 == 1)
            for ci, (_, key, fmt, _) in enumerate(columns):
                dat(ws2, ri, current_col + ci, run.get(key), fmt=fmt, is_odd=odd)
        current_col += ncols + 1

    ws2.freeze_panes = "A3"

    # ── RunLog sheet ──────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("RunLog")
    ws3.sheet_view.showGridLines = False
    title_row(ws3, "Run Log — All Saved Experiments", 4)
    for ci, label in enumerate(["Filename", "Samples", "Status", "Saved"], 1):
        hdr(ws3, 2, ci, label)
    ws3.column_dimensions["A"].width = 42
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 10
    ws3.column_dimensions["D"].width = 22
    all_jsons = sorted(
        [f for f in os.listdir(results_dir) if f.startswith("run_")],
        reverse=True,
    )
    for ri, fname in enumerate(all_jsons, 3):
        odd    = (ri % 2 == 1)
        parts  = fname.replace(".json", "").split("_")
        samps  = parts[-1].replace("samples", "") if parts else "?"
        dat(ws3, ri, 1, fname,   is_odd=odd)
        dat(ws3, ri, 2, samps,   is_odd=odd)
        dat(ws3, ri, 3, "saved", is_odd=odd)
        dat(ws3, ri, 4, ts_now,  is_odd=odd)

    out_path = os.path.join(results_dir, "benchmark_summary.xlsx")
    wb.save(out_path)
    print(f"[Benchmark] Excel saved -> {out_path}")
    return out_path


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    """Run the full benchmark suite across all sample sizes in config.yaml."""
    results_dir  = CFG["paths"]["benchmark_dir"]
    sample_sizes = CFG["benchmark"]["sample_sizes"]
    n_epochs     = CFG["benchmark"]["n_epochs_per_run"]
    port_start   = CFG["benchmark"]["port_start"]

    os.makedirs(results_dir, exist_ok=True)

    print("\n" + "="*58)
    print("  BENCHMARK SUITE")
    print(f"  Sample sizes : {sample_sizes}")
    print(f"  Epochs/run   : {n_epochs}")
    print(f"  Results dir  : {results_dir}/")
    print("="*58)

    suite_start = time.time()

    for i, n_samples in enumerate(sample_sizes):
        port   = port_start + i
        result = run_experiment(n_samples, n_epochs, port)

        save_json(result, results_dir)
        append_csv(result, results_dir)

        t    = result.get("total_time_sec", "?")
        loss = result.get("final_loss", "?")
        st   = result.get("status", "?")
        print(f"\n  -> {n_samples:,} samples | {t}s | loss={loss} | {st}")

    build_excel(results_dir)

    suite_total = round(time.time() - suite_start, 1)
    print(f"\n{'='*58}")
    print(f"  SUITE COMPLETE — {suite_total}s total")
    print(f"  Results in: {results_dir}/")
    print("="*58)


if __name__ == "__main__":
    multiprocessing.set_start_method("spawn", force=True)
    main()